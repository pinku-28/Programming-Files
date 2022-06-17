import openpyxl as xl
from openpyxl.chart import BarChart, Reference
import os

filename = input('filename: ')
worksheet = input('worksheet: ')
minrow = str(input('from row: '))
minmax_col = str(input('column: '))


def process_workbook(filename, worksheet, minrow, minmax_col):

    wb = xl.load_workbook(filename)
    sheet = wb[worksheet]

    values = Reference(sheet,
                       min_row=minrow,
                       max_row=sheet.max_row,
                       min_col=minmax_col,
                       max_col=minmax_col
                       )

    chart = BarChart()
    chart.add_data(values)
    sheet.add_chart(chart, minmax_col+1)

    wb.save(f'{filename}(charted)')


process_workbook(filename, worksheet, minrow, minmax_col)

# potek ayaw gumana edi wag
